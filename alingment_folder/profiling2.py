import json
import os
import openpyxl


aligned_names = {
    "dataset": ["dataset", "data"],
    "question": ["question type", "question combination", "question"],
    "diagram": ["diagram", "diagram #"],
    "participants": ["group", "participants"],
    "query/question type": ["query type", "question type"],
    "features": ["features", "feature"],
    "related information": ["train ex.", "annotations"],
    "checkpoints": ["used in prediction", "checkpoint"],
    "setting": ["training_setting", "setting"]
}

aligned_values = {
    "bing": ["bing", "bing full", "bing subset"],
    "segmentation": ["image segmentation", "semantic segmentation"],
    "model/checkpoint": ["1 model, 1 checkpoint", "1 model, 3 checkpoints", "2 models 3 checkpoints", "1 model, 1 checkpoint, ema", "2 models 3 checkpoints, ema", "final ensemble 3 model, 3 checkpoints, ema"],
    "fold/checkpoint": ["1 fold, 1 checkpoint", "1 fold, 3 checkpoints", "5 folds, 1 checkpoint", "5 folds, 3 checkpoints", "final ensemble 5 folds, 5 checkpoints"],
    "a1": ["(a1)", "g<sub>a1</sub>", "a1"],
    "a2": ["(a2)", "g<sub>a2</sub>", "a2"],
    "a3": ["(a3)", "g<sub>a3</sub>", "a3"],
    "b1": ["(b1)", "g<sub>b1</sub>", "b1'"],
    "b2": ["(b2)", "g<sub>b2</sub>", "b2'", "b2''"],
    "b3": ["(b3)", "g<sub>b3</sub>", "b3"],
    "b4": ["(b4)", "g<sub>b4</sub>", "b4'", "b4''"],
    "c1": ["(c1)", "g<sub>c1</sub>", "c1'", "c1''", "c1'''"],
    "c2": ["(c2)", "g<sub>c2</sub>", "c2'", "c2''", "c2'''"],
    "d1": ["(d1)", "g<sub>d1</sub>", "d1"],
    "d2": ["(d2)", "g<sub>d2</sub>", "d2"],
    "d3": ["(d3)", "g<sub>d3</sub>", "d3'", "d3''"],
    "ckpt" : ["ckpt 36050", "ckpt 48050", "ckpt 54050", "ckpt 60050", "ckpt 66050"],
    "n2n" : ["n2n", "n2n-oov"]
}

aligned_measure = {
    "avg": ["avg. rel. err.", "avg. cl. acc.", "avg"],
    "accuracy": ["accuracy", "top-5 accuracy", "acc."],
    "sqrt": ["sqrt(s1n)", "sqrt(s1c)", "sqrt(s2n)", "sqrt(s2c)"],
    "delta": ["delta_s1_max", "delta_s2_max"],
    "lb": ["private lb", "public lb"],
    "p_t": ["ρ<sub>t</sub>"]
}

def count_specifications_and_measure(json_dir, excel_file):
    if os.path.exists(excel_file):
        workbook = openpyxl.load_workbook(excel_file)
    else:
        workbook = openpyxl.Workbook()

    
    sheet_names = ["Aligned Names", "Aligned Values", "Aligned Measures"]

    for sheet_name in sheet_names:
        if sheet_name in workbook.sheetnames:
            sheet = workbook[sheet_name]
            sheet.delete_rows(1, sheet.max_row)  
        else:
            workbook.create_sheet(title=sheet_name)

    
    aligned_name_sheet = workbook[sheet_names[0]]
    aligned_value_sheet = workbook[sheet_names[1]]
    aligned_measure_sheet = workbook[sheet_names[2]]

    
    name_counts = {}
    value_counts = {}
    measure_counts = {}

    
    for file_name in os.listdir(json_dir):
        if file_name.endswith(".json"):
            with open(os.path.join(json_dir, file_name), "r") as file:
                claims = json.load(file)

                for claim_id, claim in claims.items():
                    for spec_id, spec in claim.get("Specifications", {}).items():
                        if "name" in spec:
                            found = False
                            for aligned_name, synonyms in aligned_names.items():
                                if spec["name"].lower() in synonyms:
                                    name = aligned_name
                                    name_counts[name] = name_counts.get(name, 0) + 1
                                    found = True
                                    break
                            if not found:  
                                name = spec["name"].lower()
                                name_counts[name] = name_counts.get(name, 0) + 1

                        
                        if "value" in spec:
                            found = False
                            for aligned_value, synonyms in aligned_values.items():
                                if spec["value"].lower() in synonyms:
                                    value = aligned_value
                                    value_counts[value] = value_counts.get(value, 0) + 1
                                    found = True
                                    break
                            if not found:  
                                value = spec["value"].lower()
                                value_counts[value] = value_counts.get(value, 0) + 1

                    
                    if "Measure" in claim:
                        found = False
                        for aligned_measure_name, synonyms in aligned_measure.items():
                            if claim["Measure"].lower() in synonyms:
                                measure = aligned_measure_name
                                measure_counts[measure] = measure_counts.get(measure, 0) + 1
                                found = True
                                break
                        if not found:  
                            measure = claim["Measure"].lower()
                            measure_counts[measure] = measure_counts.get(measure, 0) + 1

    
    def write_to_sheet(sheet, data):
        sheet.cell(row=1, column=1, value="Category")
        sheet.cell(row=1, column=2, value="Count")
        row = 2
        for key, count in data.items():
            sheet.cell(row=row, column=1, value=key)
            sheet.cell(row=row, column=2, value=count)
            row += 1

    write_to_sheet(aligned_name_sheet, name_counts)
    write_to_sheet(aligned_value_sheet, value_counts)
    write_to_sheet(aligned_measure_sheet, measure_counts)

    # Guardar el archivo Excel
    workbook.save(excel_file)


json_dir = r"C:\Users\Dell XPS 9510\Downloads\Homework4-fabio_branch\Homework4-fabio_branch\fabio_misino_ortega_claims"
excel_file = r"C:\Users\Dell XPS 9510\Desktop\java\t2\homework4\alingment_folder\profiling_after_alignment.xlsx"


count_specifications_and_measure(json_dir, excel_file)
